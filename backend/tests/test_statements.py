"""Synthetic workbooks that copy each bank's real layout (all names/numbers are made up)."""
import base64
import io
from datetime import datetime

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.statements import StatementError, detect_and_parse

IBAN = "AZ21NABZ00000000137010001944"


def xlsx(build) -> bytes:
    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def abb(wb):
    ws = wb.active
    ws["B26"], ws["F26"], ws["H26"] = "Çıxarışın tarixi / Date:", "Müddətin əvvəlinə qalıq / Starting balance:", "100.00 ₼"
    ws["B30"], ws["D30"] = "VÖEN / TAX ID:", "1111111111"
    ws["B32"], ws["D32"] = "IBAN / IBAN:", IBAN
    ws["B34"], ws["D34"] = "Hesabın valyutası / Account currency", "AZN"
    ws["F34"], ws["H34"] = "Yekun balans / Balance:", "145.00 ₼"
    ws.append([])
    ws["B42"], ws["C42"], ws["D42"] = "Əməliyyatın tarixi / Operation date", "Əməliyyatın referensi / Reference İD", "Əks hesab üzrə məlumatlar / Counterparty account details"
    ws["E42"], ws["F42"], ws["G42"], ws["H42"] = "Debit məbləğ / Debit amount", "Kredit məbləğ / Credit amount", "Ödənişin təyinatı / Description", "Qarşı tərəfin VÖEN-i / Counterparty TAX ID"
    ws["B43"], ws["C43"], ws["E43"], ws["F43"], ws["G43"] = "01.05.2026", "R1", 0, 100, "TRANSFER"
    ws["B44"], ws["C44"], ws["D44"] = "02.05.2026", "R2", f"/{IBAN} TEST MMC 2222222222/ 505141/"
    ws["E44"], ws["F44"], ws["G44"], ws["H44"] = 55, 0, "ödəniş", "2222222222"


def accessbank(wb):
    s = wb.active
    s.title = "Summary"
    s.append(["", "Hesab nömrəsi\nAccount No", IBAN])
    s.append(["", "Valyuta\nCurrency", "AZN"])
    s.append(["", "İlkin Balans\nStarting Balance", 100])
    s.append(["", "Son Balans\nClosing Balance", 145])
    t = wb.create_sheet("account_statement")
    t.append(["", "Hesablaşma tarixi\nValue Date", "Əməliyyat tarixi\nPosting Date", "Əməliyyatın №-si\nTransaction ID",
              "Mədaxil\nIn", "Məxaric\nOut", "Balans\nBalance", "Göndərənin/Alanın Adı\nSender / Beneficiary Name",
              "Göndərən/Alanin Tax ID\nSender/Beneficiary TAX ID", "Təyinat\nDescription"])
    d = datetime(2026, 6, 1)
    t.append(["", d, d, "FT1", 100, None, 200, f"TEST MMC\n{IBAN}\nBank", "2222222222", "Mal\nsatışı"])
    t.append(["", d, d, "FT2", None, 55, 145, None, "2222222222", "komissiya"])


def aiib(wb):
    ws = wb.active
    ws.append(["Hesabdan çıxarış / Account statement"])
    ws.append(["Hesab nömrəsi / Account number", IBAN, "Hesabın valyutası / Account currency", "AZN"])
    ws.append(["VÖEN / Tax id", "1111111111", "x", "y"])
    ws.append(["İlkin qalıq", "100 AZN", "Son qalıq", "145 AZN"])
    ws.append(["№", "Tarix", "Sənədin nömrəsi", "Əks hesabın nömrəsi, adı və vöen-i", "Məbləğ", "Valyuta", "AZN dəyəri", "Balans", "Təyinat"])
    ws.append([1, datetime(2025, 1, 6), "1 / 2", f"{IBAN} / A / B MMC / 2222222222", 100, "AZN", 100, 200, "gəlir"])
    ws.append([2, datetime(2025, 1, 6), "3 / 4", f"{IBAN} / TEST MMC / 2222222222", -55, "AZN", -55, 145, "xərc"])


def respublika(wb):
    ws = wb.active
    ws.append(["Hesab üzrə çıxarış"])
    ws.append(["Hesabın sahibi", "TEST MMC", "Giriş qalığı", 100])
    ws.append(["VÖEN", "1111111111", "Kredit dövriyyəsi", 100])
    ws.append(["Hesabın nömrəsi", IBAN, "Debet dövriyyəsi", 55])
    ws.append(["Valyutanın kodu", "AZN", "Çıxış qalığı", 145])
    ws.append(["Tarix", "Sənədin N", "Göndərənin adı/Alanın adı", "VÖEN", "Hesab", "Ödənişin təyinatı", "Kredit", "Debet", "Qalıq"])
    ws.append(["01.06.2026", "X1", "TEST MMC", "2222222222", IBAN, "gəlir", 100.0, 0.0, 200.0])
    ws.append(["02.06.2026", "X2", "TEST MMC", "2222222222", IBAN, "xərc", 0.0, 55.0, 145.0])


@pytest.mark.parametrize("build,bank", [(abb, "ABB"), (accessbank, "Accessbank"), (aiib, "AIIB"), (respublika, "Bank Respublika")])
def test_parsers(build, bank):
    st = detect_and_parse(xlsx(build))
    assert st.bank == bank and st.iban == IBAN and len(st.txs) == 2
    assert st.check() == ""
    a, b = st.to_rows()
    assert (a["c9"], a["c10"]) == (100, 0) and (b["c9"], b["c10"]) == (0, 55)
    assert (a["c12"], a["c13"], b["c13"]) == (100, 200, 145)
    assert a["c4"].endswith("T00:00:00.000") and a["c1"] == bank


def test_counterparty_split_abb():
    b = detect_and_parse(xlsx(abb)).txs[1]
    assert (b.cp_account, b.cp_name, b.cp_tin) == (IBAN, "TEST MMC", "2222222222")


def test_unknown_format():
    with pytest.raises(StatementError):
        detect_and_parse(xlsx(lambda wb: wb.active.append(["salam"])))


def test_endpoint():
    c = TestClient(app)
    r = c.post("/api/1c/statement", json={"data": base64.b64encode(xlsx(respublika)).decode()}).json()
    assert r["success"] and len(r["data"]) == 2 and "Bank Respublika" in r["message"]
    r = c.post("/api/1c/statement", json={"data": "!!!"}).json()
    assert not r["success"] and r["message"]
